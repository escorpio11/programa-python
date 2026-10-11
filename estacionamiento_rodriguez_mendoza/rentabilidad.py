"""
Presupuesto referencial (CAPEX), costos de operación (OPEX) y análisis de
rentabilidad del estacionamiento techado. Genera además una hoja de cálculo
con fórmulas vivas para que el propietario ajuste tarifas, ocupación y precios.

Montos en soles (S/), precios referenciales oct-2026 para Mendoza (Amazonas),
con IGV. Son estimaciones de anteproyecto: deben reemplazarse con cotizaciones
locales antes de decidir.
"""
from __future__ import annotations

import math
from dataclasses import dataclass, field

import estructura as est
import geometria as g

# --------------------------------------------------------------------------
# Cantidades derivadas del diseño
# --------------------------------------------------------------------------
R = g.resumen()
N_PLAZAS = R["plazas_total"]
N_MOTOS = R["motos"] + R["mototaxis"]
AREA_LOTE = R["area_lote"]
AREA_PAVIMENTO = round(R["area_circulacion"] + R["area_plazas"] + sum(m.area for m in g.MOTOS + g.MOTOTAXIS)
                       + g.PASO_MOTO.area, 1)
ML_CERCO_FRENTES = round(g.LOTE.exterior.length * 0 + 30.0 + 20.0 - (3.0 + 3.0 + 1.0 + 1.2), 1)
ML_CERCO_FONDO = 30.0
ML_CANALETA = round(sum(c.s1 - c.s0 for c in g.COBERTURAS), 1)
N_COLUMNAS = len(g.COLUMNAS)
AREA_TECHO = round(est.AREA_TECHADA, 1)
KG_ACERO = round(est.ACERO_TOTAL_KG, 0)
VOL_CONCRETO_ZAPATA = est.ZAPATA["vol_zap"] + est.ZAPATA["vol_dado"]
N_LUMINARIAS = math.ceil(N_PLAZAS / 2)


@dataclass
class Partida:
    grupo: str
    descripcion: str
    unidad: str
    cantidad: float
    precio: float
    clave: str = ""

    @property
    def parcial(self):
        return self.cantidad * self.precio


def costo_zapata():
    """Costo unitario de una zapata Z-1 con dado y pernos (S/)."""
    exc = est.ZAP_B * est.ZAP_L * (est.DF_ZAPATA + 0.10) * 1.25
    concreto = VOL_CONCRETO_ZAPATA
    solado = est.ZAP_B * est.ZAP_L * 0.10
    acero_kg = 72.0
    return round(exc * 45 + concreto * 560 + solado * 320 + acero_kg * 6.2 + 1.6 * 48 + 260, 0)


def partidas(piso="afirmado"):
    pz = costo_zapata()
    piso_pu = 40.0 if piso == "afirmado" else 130.0
    piso_txt = ("Afirmado compactado e=0.20 m (95 % PM) con bombeo" if piso == "afirmado"
                else "Losa de concreto f'c=210 e=0.15 m sobre base granular e=0.15 m")
    return [
        Partida("Obras preliminares", "Limpieza, retiro de capa orgánica e=0.20 m y eliminación", "m²", round(AREA_LOTE, 1), 8.0, "prelim"),
        Partida("Obras preliminares", "Trazo, nivelación y replanteo", "glb", 1, 1500.0),
        Partida("Pavimento", piso_txt, "m²", AREA_PAVIMENTO, piso_pu, "piso"),
        Partida("Cerco y accesos", "Cerco de albañilería confinada h=2.40 m (frentes a calle)", "ml", ML_CERCO_FRENTES, 330.0),
        Partida("Cerco y accesos", "Cerco de albañilería h=2.40 m (fondo C-D)", "ml", ML_CERCO_FONDO, 330.0),
        Partida("Cerco y accesos", "Portón metálico corredizo 3.00 x 2.40 m", "und", 2, 4800.0),
        Partida("Cerco y accesos", "Puerta peatonal metálica 1.00 x 2.10 m", "und", 2, 950.0),
        Partida("Estructura techada", "Zapata Z-1 + dado + pernos de anclaje", "und", N_COLUMNAS, pz),
        Partida("Estructura techada", "Estructura metálica: suministro, fabricación, pintura y montaje", "kg", KG_ACERO, 12.0, "acero"),
        Partida("Estructura techada", "Cobertura TR-4 0.40 mm prepintada + fascia y accesorios", "m²", round(AREA_TECHO * 1.06, 1), 55.0),
        Partida("Estructura techada", "Canaleta galvanizada e=0.6 mm", "ml", ML_CANALETA, 60.0),
        Partida("Estructura techada", "Bajadas pluviales PVC ø 4\"", "und", N_COLUMNAS, 220.0),
        Partida("Drenaje", "Canaleta de piso con rejilla metálica y descarga a vía", "ml", 40, 220.0),
        Partida("Instalaciones eléctricas", "Acometida, medidor y tablero general", "glb", 1, 3800.0),
        Partida("Instalaciones eléctricas", "Luminaria LED hermética 40 W bajo techo", "und", N_LUMINARIAS, 260.0),
        Partida("Instalaciones eléctricas", "Poste metálico 6 m con reflector LED 100 W", "und", 3, 1800.0),
        Partida("Instalaciones eléctricas", "Cableado, tuberías y tomacorrientes", "glb", 1, 3000.0),
        Partida("Instalaciones eléctricas", "Sistema CCTV (6 cámaras + grabador)", "glb", 1, 4200.0),
        Partida("Servicios", "Garita, SS.HH. accesible y depósito (albañilería)", "m²",
                round(g.GARITA.area + g.SSHH.area + g.DEPOSITO.area, 1), 1500.0),
        Partida("Servicios", "Instalaciones sanitarias (agua, desagüe y conexiones)", "glb", 1, 3800.0),
        Partida("Señalización y seguridad", "Pintura de tráfico: líneas, numeración y flechas", "glb", 1, 2800.0),
        Partida("Señalización y seguridad", "Topellantas de concreto prefabricado", "und", N_PLAZAS, 75.0),
        Partida("Señalización y seguridad", "Señales verticales, espejo convexo y bolardos", "glb", 1, 2050.0),
        Partida("Señalización y seguridad", "Extintores PQS 6 kg y señalética de evacuación", "glb", 1, 870.0),
        Partida("Estudios y licencias", "Topografía, EMS, expediente técnico firmado y licencias (edificación, funcionamiento, ITSE)", "glb", 1, 13000.0),
    ]


GASTOS_GENERALES = 0.10      # dirección técnica / supervisión / utilidad del contratista
IMPREVISTOS = 0.07


def capex(piso="afirmado"):
    ps = partidas(piso)
    sub = sum(p.parcial for p in ps)
    gg = sub * GASTOS_GENERALES
    imp = (sub + gg) * IMPREVISTOS
    return ps, sub, gg, imp, sub + gg + imp


def capex_por_grupo(piso="afirmado"):
    ps, sub, gg, imp, tot = capex(piso)
    grupos = {}
    for p in ps:
        grupos[p.grupo] = grupos.get(p.grupo, 0.0) + p.parcial
    grupos["Gastos generales e imprevistos"] = gg + imp
    return grupos, tot


# --------------------------------------------------------------------------
# Operación e ingresos
# --------------------------------------------------------------------------
@dataclass
class Escenario:
    nombre: str
    abonados: int            # autos con abono mensual (principalmente nocturno)
    tarifa_mensual: float    # S/ por mes
    tickets_dia: float       # autos de rotación por día (horas / día / noche)
    ticket_prom: float       # S/ por ticket
    motos: int               # motos y mototaxis abonados
    tarifa_moto: float


ESCENARIOS = [
    Escenario("Pesimista", 12, 100.0, 8, 4.0, 4, 35.0),
    Escenario("Base", 18, 120.0, 12, 4.5, 6, 40.0),
    Escenario("Optimista", 24, 140.0, 18, 5.0, 7, 45.0),
]

SUELDO_RMV = 1130.0          # remuneración mínima vital (verificar la vigente)
FACTOR_LABORAL = 1.09 + 1 / 24   # EsSalud 9 % + vacaciones régimen microempresa
TARIFA_KWH = 0.80
KWH_DIA = (N_LUMINARIAS * 0.040 + 3 * 0.100) * 11 + 0.060 * 24
AGUA_MES = 35.0
MANTENIMIENTO = 0.010        # del CAPEX por año
PREDIAL_ARBITRIOS = 1200.0
SEGURO = 1800.0
NRUS_MES = 50.0              # Nuevo RUS categoría 2 (ingresos <= S/ 8,000/mes)
VARIOS_MES = 100.0           # tickets, internet, útiles
RAMPA = [0.60, 0.85]         # ocupación relativa en años 1 y 2
CRECIMIENTO = 0.03           # ajuste anual de tarifas y costos
TASA = 0.10                  # tasa de descuento anual
HORIZONTE = 20
MANT_MAYOR_ANIO = 10
MANT_MAYOR = 0.06            # del CAPEX (repintado de estructura, reposición de afirmado)


GRUPO_TECHO = "Estructura techada"


def capex_sin_techo(piso="afirmado"):
    """Inversión si se construye primero sin coberturas (fase 1)."""
    ps = [p for p in partidas(piso) if p.grupo != GRUPO_TECHO]
    sub = sum(p.parcial for p in ps)
    gg = sub * GASTOS_GENERALES
    imp = (sub + gg) * IMPREVISTOS
    return sub + gg + imp


# Con plazas descubiertas en una zona muy lluviosa se asume menor tarifa y ocupación
ESCENARIO_SIN_TECHO = Escenario("Base sin techo", 15, 95.0, 11, 4.0, 6, 40.0)


def ingreso_mensual(e: Escenario):
    return e.abonados * e.tarifa_mensual + e.tickets_dia * e.ticket_prom * 30 + e.motos * e.tarifa_moto


def opex_anual(capex_total, con_personal=True):
    personal = SUELDO_RMV * FACTOR_LABORAL * 12 if con_personal else 0.0
    energia = KWH_DIA * 30 * TARIFA_KWH * 12
    return {
        "Personal (1 guardián-cobrador)": personal,
        "Energía eléctrica": energia,
        "Agua y desagüe": AGUA_MES * 12,
        "Mantenimiento (1 % del CAPEX)": MANTENIMIENTO * capex_total,
        "Predial y arbitrios": PREDIAL_ARBITRIOS,
        "Seguro multirriesgo y RC": SEGURO,
        "Tributos (Nuevo RUS)": NRUS_MES * 12,
        "Varios (tickets, internet, útiles)": VARIOS_MES * 12,
    }


def flujo(e: Escenario, capex_total, con_personal=True):
    filas = []
    acum = -capex_total
    filas.append(dict(anio=0, ingresos=0, opex=0, otros=0, neto=-capex_total, acum=acum))
    base_opex = sum(opex_anual(capex_total, con_personal).values())
    for t in range(1, HORIZONTE + 1):
        f = RAMPA[t - 1] if t <= len(RAMPA) else 1.0
        crec = (1 + CRECIMIENTO) ** (t - 1)
        ing = ingreso_mensual(e) * 12 * f * crec
        op = base_opex * crec
        otros = MANT_MAYOR * capex_total if t == MANT_MAYOR_ANIO else 0.0
        neto = ing - op - otros
        acum += neto
        filas.append(dict(anio=t, ingresos=ing, opex=op, otros=otros, neto=neto, acum=acum))
    return filas


def van(flujos, tasa=TASA):
    return sum(f / (1 + tasa) ** t for t, f in enumerate(flujos))


def tir(flujos):
    lo, hi = -0.99, 1.0
    f_lo, f_hi = van(flujos, lo), van(flujos, hi)
    if f_lo * f_hi > 0:
        return None
    for _ in range(200):
        mid = (lo + hi) / 2
        fm = van(flujos, mid)
        if f_lo * fm <= 0:
            hi, f_hi = mid, fm
        else:
            lo, f_lo = mid, fm
    return (lo + hi) / 2


def recuperacion(filas):
    for a, b in zip(filas, filas[1:]):
        if a["acum"] < 0 <= b["acum"]:
            return a["anio"] + (-a["acum"]) / b["neto"]
    return None


def evaluar(piso="afirmado", con_personal=True):
    *_, total = capex(piso)
    res = []
    for e in ESCENARIOS:
        fl = flujo(e, total, con_personal)
        netos = [f["neto"] for f in fl]
        op = sum(opex_anual(total, con_personal).values())
        ing = ingreso_mensual(e) * 12
        res.append(dict(escenario=e, capex=total, ingreso_mes=ingreso_mensual(e), ingreso_anual=ing, opex=op,
                        ebitda=ing - op, van=van(netos), tir=tir(netos), payback=recuperacion(fl), flujo=fl))
    return res


def punto_equilibrio(capex_total, con_personal=True, anios=10, base=ESCENARIOS[1]):
    """Abonados necesarios (con la rotación y motos del escenario base) para:
    a) cubrir el OPEX; b) recuperar la inversión en `anios` años (sin crecimiento)."""
    op = sum(opex_anual(capex_total, con_personal).values())
    otros = (base.tickets_dia * base.ticket_prom * 30 + base.motos * base.tarifa_moto) * 12
    a = max(0.0, (op - otros) / (base.tarifa_mensual * 12))
    b = max(0.0, (op + capex_total / anios - otros) / (base.tarifa_mensual * 12))
    return a, b


# --------------------------------------------------------------------------
# Hoja de cálculo con fórmulas
# --------------------------------------------------------------------------
def exportar_xlsx(ruta):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    azul = PatternFill("solid", fgColor="DCE6F1")
    gris = PatternFill("solid", fgColor="EDEDED")
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    negrita = Font(bold=True)
    titulo = Font(bold=True, size=14)
    fino = Side(style="thin", color="999999")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    soles = '"S/ "#,##0'
    soles2 = '"S/ "#,##0.00'

    # ---------------- Supuestos
    ws = wb.active
    ws.title = "Supuestos"
    ws["A1"] = "SUPUESTOS DEL MODELO (celdas amarillas = editables)"
    ws["A1"].font = titulo
    sup = [
        ("Datos del proyecto", None, None, None),
        ("Plazas vehiculares techadas", N_PLAZAS, "und", "plazas"),
        ("Plazas motos / mototaxis", N_MOTOS, "und", "motos_cap"),
        ("Área del terreno", round(AREA_LOTE, 1), "m²", "area_lote"),
        ("Área pavimentada", AREA_PAVIMENTO, "m²", "area_piso"),
        ("Área techada", AREA_TECHO, "m²", "area_techo"),
        ("Acero estructural", KG_ACERO, "kg", "kg_acero"),
        ("Tipo de piso (1 = afirmado, 2 = losa de concreto)", 1, "", "tipo_piso"),
        ("Precio afirmado e=0.20 m", 40.0, "S/ por m²", "pu_afirmado"),
        ("Precio losa de concreto e=0.15 m", 130.0, "S/ por m²", "pu_losa"),
        ("Precio estructura metálica instalada", 12.0, "S/ por kg", "pu_acero"),
        ("Gastos generales y dirección técnica", GASTOS_GENERALES, "%", "gg"),
        ("Imprevistos", IMPREVISTOS, "%", "imprev"),
        ("Operación", None, None, None),
        ("¿Guardián contratado? (1 = sí, 0 = lo opera el propietario)", 1, "", "personal"),
        ("Remuneración mínima vital", SUELDO_RMV, "S/ por mes", "rmv"),
        ("Factor de cargas laborales (EsSalud + vacaciones)", round(FACTOR_LABORAL, 4), "", "f_lab"),
        ("Consumo eléctrico", round(KWH_DIA, 2), "kWh por día", "kwh"),
        ("Tarifa eléctrica", TARIFA_KWH, "S/ por kWh", "t_kwh"),
        ("Agua y desagüe", AGUA_MES, "S/ por mes", "agua"),
        ("Mantenimiento anual", MANTENIMIENTO, "% del CAPEX", "mant"),
        ("Predial y arbitrios", PREDIAL_ARBITRIOS, "S/ por año", "predial"),
        ("Seguro multirriesgo y RC", SEGURO, "S/ por año", "seguro"),
        ("Cuota Nuevo RUS", NRUS_MES, "S/ por mes", "nrus"),
        ("Varios (tickets, internet)", VARIOS_MES, "S/ por mes", "varios"),
        ("Evaluación", None, None, None),
        ("Ocupación relativa año 1", RAMPA[0], "%", "rampa1"),
        ("Ocupación relativa año 2", RAMPA[1], "%", "rampa2"),
        ("Crecimiento anual de tarifas y costos", CRECIMIENTO, "%", "crec"),
        ("Tasa de descuento", TASA, "%", "tasa"),
        ("Mantenimiento mayor (año 10)", MANT_MAYOR, "% del CAPEX", "mant_mayor"),
    ]
    ref = {}
    fila = 3
    for etiqueta, valor, unidad, clave in sup:
        if valor is None:
            ws.cell(fila, 1, etiqueta).font = negrita
            ws.cell(fila, 1).fill = gris
            fila += 1
            continue
        ws.cell(fila, 1, etiqueta)
        c = ws.cell(fila, 2, valor)
        c.fill = amarillo
        c.border = borde
        if unidad in ("%", "% del CAPEX"):
            c.number_format = "0.0%"
        ws.cell(fila, 3, unidad)
        ref[clave] = f"Supuestos!$B${fila}"
        fila += 1
    dv = DataValidation(type="list", formula1='"1,2"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ref["tipo_piso"].split("!")[1].replace("$", ""))
    dv2 = DataValidation(type="list", formula1='"0,1"', allow_blank=False)
    ws.add_data_validation(dv2)
    dv2.add(ref["personal"].split("!")[1].replace("$", ""))
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14

    # ---------------- Inversión
    wi = wb.create_sheet("Inversion")
    wi["A1"] = "PRESUPUESTO REFERENCIAL DE INVERSIÓN (CAPEX) - S/ con IGV"
    wi["A1"].font = titulo
    cab = ["Grupo", "Partida", "Und.", "Cantidad", "P. unitario", "Parcial"]
    for j, t in enumerate(cab, 1):
        c = wi.cell(3, j, t)
        c.font = negrita
        c.fill = azul
        c.border = borde
    r = 4
    primera = r
    for p in partidas("afirmado"):
        wi.cell(r, 1, p.grupo)
        wi.cell(r, 2, p.descripcion if p.clave != "piso" else "Pavimento (afirmado o losa según Supuestos)")
        wi.cell(r, 3, p.unidad)
        if p.clave == "piso":
            wi.cell(r, 4, f"={ref['area_piso']}")
            wi.cell(r, 5, f"=IF({ref['tipo_piso']}=1,{ref['pu_afirmado']},{ref['pu_losa']})")
        elif p.clave == "acero":
            wi.cell(r, 4, f"={ref['kg_acero']}")
            wi.cell(r, 5, f"={ref['pu_acero']}")
        else:
            wi.cell(r, 4, p.cantidad)
            c = wi.cell(r, 5, p.precio)
            c.fill = amarillo
        wi.cell(r, 6, f"=D{r}*E{r}")
        wi.cell(r, 5).number_format = soles2
        wi.cell(r, 6).number_format = soles
        for j in range(1, 7):
            wi.cell(r, j).border = borde
        r += 1
    ultima = r - 1
    wi.cell(r, 5, "Subtotal").font = negrita
    wi.cell(r, 6, f"=SUM(F{primera}:F{ultima})").number_format = soles
    r_sub = r
    r += 1
    wi.cell(r, 5, "Gastos generales")
    wi.cell(r, 6, f"=F{r_sub}*{ref['gg']}").number_format = soles
    r_gg = r
    r += 1
    wi.cell(r, 5, "Imprevistos")
    wi.cell(r, 6, f"=(F{r_sub}+F{r_gg})*{ref['imprev']}").number_format = soles
    r_imp = r
    r += 1
    wi.cell(r, 5, "TOTAL INVERSIÓN").font = negrita
    c = wi.cell(r, 6, f"=F{r_sub}+F{r_gg}+F{r_imp}")
    c.number_format = soles
    c.font = negrita
    ref["capex"] = f"Inversion!$F${r}"
    r += 1
    wi.cell(r, 5, "Por plaza techada")
    wi.cell(r, 6, f"={ref['capex']}/{ref['plazas']}").number_format = soles
    for col, w in zip("ABCDEF", (24, 70, 7, 11, 14, 14)):
        wi.column_dimensions[col].width = w

    # ---------------- Escenarios
    we = wb.create_sheet("Escenarios")
    we["A1"] = "ESCENARIOS DE DEMANDA Y RESULTADOS (año de operación estable)"
    we["A1"].font = titulo
    for j, e in enumerate(ESCENARIOS):
        c = we.cell(3, 2 + j, e.nombre)
        c.font = negrita
        c.fill = azul
        c.alignment = Alignment(horizontal="center")
    filas_e = [
        ("Autos abonados mensuales (nocturno)", "abonados", "n"),
        ("Tarifa mensual por auto (S/)", "tarifa_mensual", "s"),
        ("Autos de rotación por día", "tickets_dia", "n"),
        ("Ticket promedio de rotación (S/)", "ticket_prom", "s"),
        ("Motos / mototaxis abonados", "motos", "n"),
        ("Tarifa mensual moto (S/)", "tarifa_moto", "s"),
    ]
    fe = {}
    r = 4
    for etiqueta, attr, tipo in filas_e:
        we.cell(r, 1, etiqueta)
        for j, e in enumerate(ESCENARIOS):
            c = we.cell(r, 2 + j, getattr(e, attr))
            c.fill = amarillo
            c.border = borde
            if tipo == "s":
                c.number_format = soles2
        fe[attr] = r
        r += 1
    r += 1
    we.cell(r, 1, "Ingreso mensual").font = negrita
    for j in range(3):
        col = get_column_letter(2 + j)
        we.cell(r, 2 + j, f"={col}{fe['abonados']}*{col}{fe['tarifa_mensual']}+{col}{fe['tickets_dia']}*{col}{fe['ticket_prom']}*30+{col}{fe['motos']}*{col}{fe['tarifa_moto']}").number_format = soles
    fe["ing_mes"] = r
    r += 1
    we.cell(r, 1, "Ingreso anual")
    for j in range(3):
        col = get_column_letter(2 + j)
        we.cell(r, 2 + j, f"={col}{fe['ing_mes']}*12").number_format = soles
    fe["ing_anual"] = r
    r += 2
    we.cell(r, 1, "Costos de operación anuales").font = negrita
    r += 1
    opex_items = [
        ("Personal (1 guardián-cobrador)", f"={ref['personal']}*{ref['rmv']}*{ref['f_lab']}*12"),
        ("Energía eléctrica", f"={ref['kwh']}*30*{ref['t_kwh']}*12"),
        ("Agua y desagüe", f"={ref['agua']}*12"),
        ("Mantenimiento", f"={ref['mant']}*{ref['capex']}"),
        ("Predial y arbitrios", f"={ref['predial']}"),
        ("Seguro", f"={ref['seguro']}"),
        ("Tributos (Nuevo RUS)", f"={ref['nrus']}*12"),
        ("Varios", f"={ref['varios']}*12"),
    ]
    r_op0 = r
    for etiqueta, f in opex_items:
        we.cell(r, 1, "   " + etiqueta)
        for j in range(3):
            we.cell(r, 2 + j, f).number_format = soles
        r += 1
    we.cell(r, 1, "Total OPEX anual").font = negrita
    for j in range(3):
        col = get_column_letter(2 + j)
        we.cell(r, 2 + j, f"=SUM({col}{r_op0}:{col}{r - 1})").number_format = soles
    fe["opex"] = r
    r += 2
    we.cell(r, 1, "Utilidad operativa anual (EBITDA)").font = negrita
    for j in range(3):
        col = get_column_letter(2 + j)
        c = we.cell(r, 2 + j, f"={col}{fe['ing_anual']}-{col}{fe['opex']}")
        c.number_format = soles
        c.font = negrita
    fe["ebitda"] = r
    r += 1
    we.cell(r, 1, "Inversión total (CAPEX)")
    for j in range(3):
        we.cell(r, 2 + j, f"={ref['capex']}").number_format = soles
    fe["capex"] = r
    r += 1
    we.cell(r, 1, "Recuperación simple (años)")
    for j in range(3):
        col = get_column_letter(2 + j)
        we.cell(r, 2 + j, f'=IF({col}{fe["ebitda"]}<=0,"No recupera",IF({col}{fe["capex"]}/{col}{fe["ebitda"]}>30,"Más de 30",{col}{fe["capex"]}/{col}{fe["ebitda"]}))').number_format = "0.0"
    r += 1
    fe["van"] = r
    fe["tir"] = r + 1
    we.cell(r, 1, f"VAN a {HORIZONTE} años (tasa en Supuestos)")
    we.cell(r + 1, 1, f"TIR a {HORIZONTE} años")
    we.cell(r + 2, 1, "Nota: VAN y TIR usan el flujo de la hoja 'Flujo' (rampa de ocupación, crecimiento y mantenimiento mayor).")
    we.column_dimensions["A"].width = 46
    for col in "BCD":
        we.column_dimensions[col].width = 16

    # ---------------- Flujo
    wf = wb.create_sheet("Flujo")
    wf["A1"] = f"FLUJO DE CAJA ANUAL ({HORIZONTE} años) - S/"
    wf["A1"].font = titulo
    wf.cell(3, 1, "Año").font = negrita
    for t in range(HORIZONTE + 1):
        wf.cell(3, 2 + t, t).font = negrita
    r = 4
    for j, e in enumerate(ESCENARIOS):
        col_e = get_column_letter(2 + j)
        wf.cell(r, 1, e.nombre).font = negrita
        wf.cell(r, 1).fill = azul
        r += 1
        r_ing, r_op, r_mm, r_neto, r_acum = r, r + 1, r + 2, r + 3, r + 4
        wf.cell(r_ing, 1, "Ingresos")
        wf.cell(r_op, 1, "OPEX")
        wf.cell(r_mm, 1, "Mantenimiento mayor")
        wf.cell(r_neto, 1, "Flujo neto")
        wf.cell(r_acum, 1, "Flujo acumulado")
        for t in range(HORIZONTE + 1):
            col = get_column_letter(2 + t)
            if t == 0:
                wf.cell(r_ing, 2, 0)
                wf.cell(r_op, 2, 0)
                wf.cell(r_mm, 2, 0)
                wf.cell(r_neto, 2, f"=-{ref['capex']}")
                wf.cell(r_acum, 2, f"=B{r_neto}")
            else:
                rampa = ref["rampa1"] if t == 1 else (ref["rampa2"] if t == 2 else "1")
                wf.cell(r_ing, 2 + t,
                        f"=Escenarios!${col_e}${fe['ing_anual']}*{rampa}*(1+{ref['crec']})^({t}-1)")
                wf.cell(r_op, 2 + t, f"=Escenarios!${col_e}${fe['opex']}*(1+{ref['crec']})^({t}-1)")
                wf.cell(r_mm, 2 + t, f"={ref['mant_mayor']}*{ref['capex']}" if t == MANT_MAYOR_ANIO else 0)
                wf.cell(r_neto, 2 + t, f"={col}{r_ing}-{col}{r_op}-{col}{r_mm}")
                prev = get_column_letter(1 + t)
                wf.cell(r_acum, 2 + t, f"={prev}{r_acum}+{col}{r_neto}")
            for rr in (r_ing, r_op, r_mm, r_neto, r_acum):
                wf.cell(rr, 2 + t).number_format = '#,##0'
        last = get_column_letter(2 + HORIZONTE)
        we.cell(fe["van"], 2 + j, f"=B{r_neto}+NPV({ref['tasa']},Flujo!C{r_neto}:{last}{r_neto})".replace(
            f"=B{r_neto}", f"=Flujo!B{r_neto}")).number_format = soles
        we.cell(fe["tir"], 2 + j, f'=IFERROR(IRR(Flujo!B{r_neto}:{last}{r_neto}),"Negativa")').number_format = "0.0%"
        r = r_acum + 2
    wf.column_dimensions["A"].width = 22

    # ---------------- Sensibilidad
    wsn = wb.create_sheet("Sensibilidad")
    wsn["A1"] = "SENSIBILIDAD: utilidad operativa anual (S/) según autos abonados y tarifa mensual"
    wsn["A1"].font = titulo
    wsn["A2"] = "Rotación, motos y OPEX del escenario Base. Años de recuperación simple en la segunda tabla."
    tarifas = [80, 100, 120, 140, 160]
    abon = [8, 12, 16, 20, 24, 28]
    base_col = "C"
    otros = (f"(Escenarios!${base_col}${fe['tickets_dia']}*Escenarios!${base_col}${fe['ticket_prom']}*30"
             f"+Escenarios!${base_col}${fe['motos']}*Escenarios!${base_col}${fe['tarifa_moto']})*12")
    for blk, (r0, titulo_blk) in enumerate(((4, "Utilidad operativa anual"), (14, "Años de recuperación"))):
        wsn.cell(r0, 1, titulo_blk).font = negrita
        wsn.cell(r0 + 1, 1, "Abonados \\ Tarifa").font = negrita
        for j, tf in enumerate(tarifas):
            c = wsn.cell(r0 + 1, 2 + j, tf)
            c.font = negrita
            c.fill = azul
            c.number_format = soles
        for i, a in enumerate(abon):
            c = wsn.cell(r0 + 2 + i, 1, a)
            c.font = negrita
            c.fill = azul
            for j in range(len(tarifas)):
                colh = get_column_letter(2 + j)
                util = f"($A{r0 + 2 + i}*{colh}${r0 + 1}*12+{otros}-Escenarios!${base_col}${fe['opex']})"
                if blk == 0:
                    f = "=" + util
                    fmt = soles
                else:
                    f = f'=IF({util}<=0,"No recupera",IF({ref["capex"]}/{util}>30,"Más de 30",{ref["capex"]}/{util}))'
                    fmt = "0.0"
                c = wsn.cell(r0 + 2 + i, 2 + j, f)
                c.number_format = fmt
                c.border = borde
    wsn.column_dimensions["A"].width = 20
    for j in range(len(tarifas)):
        wsn.column_dimensions[get_column_letter(2 + j)].width = 14
    wb.save(ruta)
    return ruta


if __name__ == "__main__":
    ps, sub, gg, imp, tot = capex()
    for p in ps:
        print(f"{p.grupo:28s} {p.descripcion[:60]:60s} {p.cantidad:>9} {p.unidad:4s} x {p.precio:>9,.2f} = {p.parcial:>11,.0f}")
    print(f"Subtotal {sub:,.0f}  GG {gg:,.0f}  Imprev {imp:,.0f}  TOTAL {tot:,.0f}  por plaza {tot / N_PLAZAS:,.0f}")
    print("Con losa:", f"{capex('losa')[-1]:,.0f}")
    print("Zapata unit:", costo_zapata())
    print("OPEX:", {k: round(v) for k, v in opex_anual(tot).items()}, round(sum(opex_anual(tot).values())))
    for cp in (True, False):
        print("Con personal" if cp else "Operado por propietario")
        for r in evaluar("afirmado", cp):
            e = r["escenario"]
            print(f"  {e.nombre:10s} ing/mes {r['ingreso_mes']:>8,.0f} EBITDA {r['ebitda']:>9,.0f} VAN {r['van']:>11,.0f} "
                  f"TIR {('%.1f%%' % (100 * r['tir'])) if r['tir'] is not None else 'n/a':>7} payback {r['payback'] if r['payback'] is None else round(r['payback'], 1)}")
    print("Equilibrio (abonados):", [round(x, 1) for x in punto_equilibrio(tot)], [round(x, 1) for x in punto_equilibrio(tot, False)])
